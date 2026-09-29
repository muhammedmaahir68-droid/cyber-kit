import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def embed_model_diagram(pptx_path, diagram_path="presentation/ncis_architecture_model_diagram.png"):
    prs = pptx.Presentation(pptx_path)

    with open(diagram_path, "rb") as f:
        diagram_bytes = f.read()

    # ──────────────────────────────────────────────────────────
    # SLIDE 2: EMBED ARCHITECTURE MODEL DIAGRAM IN COLUMN 2
    # ──────────────────────────────────────────────────────────
    s2 = prs.slides[1]

    # Find Picture shape in Slide 2
    pic_shape = None
    for shape in s2.shapes:
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE and shape.left > 2500000 and shape.left < 5800000:
            pic_shape = shape
            break

    if pic_shape:
        try:
            rId = pic_shape._element.xpath('.//a:blip/@r:embed')[0]
            image_part = s2.part.related_part(rId)
            image_part._blob = diagram_bytes
            print(f"Updated Slide 2 Picture shape image blob in {pptx_path}")

            # Re-position and resize cleanly to fit the container
            pic_shape.left = 3040000
            pic_shape.top = 1320000
            pic_shape.width = 2600000
            pic_shape.height = 1240000
        except Exception as e:
            print(f"Could not update Picture shape: {e}")

    # Remove overlapping text boxes (Shapes 18 and 19) if they still exist
    shapes_to_remove = []
    for shape in s2.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text
            if "FEATURED: Agentic AI Voice Module" in txt or "Siri-Style Multilingual Voice Agent" in txt:
                shapes_to_remove.append(shape)
            elif "Dual Core:" in txt or "Concept Diagram" in txt:
                shape.text_frame.text = "System Model: Multi-Vendor DVR + Agentic Voice AI"
                if len(shape.text_frame.paragraphs) > 0 and len(shape.text_frame.paragraphs[0].runs) > 0:
                    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)
                    shape.text_frame.paragraphs[0].runs[0].font.bold = True
                    shape.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(56, 189, 248)

    for shape in shapes_to_remove:
        try:
            sp = shape._element
            sp.getparent().remove(sp)
            print("Removed overlapping text box on Slide 2")
        except Exception as e:
            print(f"Error removing shape: {e}")

    prs.save(pptx_path)
    print(f"Model diagram successfully embedded and saved in: {pptx_path}")

if __name__ == "__main__":
    targets = [
        "presentation/SIH_Ideate_Template_NCIS.pptx",
        "presentation/SIH_Ideate_Template_AAROHAN-X.pptx",
        "frontend/public/SIH_Ideate_Template_NCIS.pptx",
        "frontend/public/SIH_Ideate_Template_AAROHAN-X.pptx"
    ]
    for t in targets:
        embed_model_diagram(t)
